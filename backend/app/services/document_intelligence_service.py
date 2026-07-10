from __future__ import annotations

import io
import logging
import os
from dataclasses import asdict, dataclass
from typing import Any, Literal

logger = logging.getLogger(__name__)

ExtractionStatus = Literal["processing", "searchable", "ocr_required", "extraction_failed"]

MIN_EXTRACTED_CHARACTERS = int(os.environ.get("OMNIX_MIN_EXTRACTED_CHARACTERS", "20"))
OCR_ENABLED = os.environ.get("OMNIX_OCR_ENABLED", "true").lower() not in {"0", "false", "no"}
MAX_OCR_PAGES = int(os.environ.get("OMNIX_MAX_OCR_PAGES", "25"))


@dataclass(slots=True)
class ExtractionDiagnostics:
    page_count: int | None = None
    extractor_used: str | None = None
    extracted_character_count: int = 0
    image_page_count: int = 0
    text_page_count: int = 0
    extraction_status: ExtractionStatus = "processing"
    extraction_failure_reason: str | None = None
    ocr_used: bool = False
    ocr_character_count: int = 0

    def to_metadata(self) -> dict[str, Any]:
        return asdict(self)


@dataclass(slots=True)
class ExtractionResult:
    text: str
    diagnostics: ExtractionDiagnostics


def normalize_extracted_text(text: str) -> str:
    return "\n\n".join(line.strip() for line in (text or "").splitlines() if line.strip())


def extraction_columns_payload(diagnostics: ExtractionDiagnostics) -> dict[str, Any]:
    return diagnostics.to_metadata()


def diagnostics_from_file(file_row: dict[str, Any]) -> ExtractionDiagnostics:
    metadata = file_row.get("metadata") if isinstance(file_row.get("metadata"), dict) else {}
    values = {**metadata}
    diagnostic_keys = set(ExtractionDiagnostics().__dataclass_fields__)
    has_any_diagnostics = any(key in metadata for key in diagnostic_keys) or any(
        key in file_row and file_row.get(key) is not None for key in diagnostic_keys
    )
    for key in diagnostic_keys:
        if key in file_row and file_row.get(key) is not None:
            values[key] = file_row.get(key)

    status = values.get("extraction_status")
    if status not in {"processing", "searchable", "ocr_required", "extraction_failed"}:
        status = "searchable" if not has_any_diagnostics or int(values.get("extracted_character_count") or 0) > 0 else "processing"

    return ExtractionDiagnostics(
        page_count=_optional_int(values.get("page_count")),
        extractor_used=_optional_str(values.get("extractor_used")),
        extracted_character_count=int(values.get("extracted_character_count") or 0),
        image_page_count=int(values.get("image_page_count") or 0),
        text_page_count=int(values.get("text_page_count") or 0),
        extraction_status=status,  # type: ignore[arg-type]
        extraction_failure_reason=_optional_str(values.get("extraction_failure_reason")),
        ocr_used=bool(values.get("ocr_used") or False),
        ocr_character_count=int(values.get("ocr_character_count") or 0),
    )


def extract_document_with_diagnostics(
    filename: str,
    file_type: str | None,
    data: bytes,
) -> ExtractionResult:
    diagnostics = ExtractionDiagnostics(extraction_status="processing")
    lowered = (file_type or "").lower()
    name_l = filename.lower()

    try:
        if "word" in lowered or name_l.endswith(".docx"):
            text = _extract_docx(data)
            diagnostics.extractor_used = "python-docx"
            diagnostics.extracted_character_count = len(normalize_extracted_text(text))
            diagnostics.text_page_count = 1 if diagnostics.extracted_character_count else 0
            diagnostics.extraction_status = "searchable" if diagnostics.extracted_character_count else "extraction_failed"
            if diagnostics.extraction_status == "extraction_failed":
                diagnostics.extraction_failure_reason = "No readable text was extracted from this DOCX file."
            return ExtractionResult(text=normalize_extracted_text(text), diagnostics=diagnostics)

        if _is_pdf(filename, file_type):
            return _extract_pdf(filename, file_type, data)

        text = normalize_extracted_text(_decode_text(data))
        diagnostics.extractor_used = "markdown" if name_l.endswith((".md", ".markdown")) or "markdown" in lowered else "text"
        diagnostics.extracted_character_count = len(text)
        diagnostics.text_page_count = 1 if text else 0
        diagnostics.extraction_status = "searchable" if diagnostics.extracted_character_count else "extraction_failed"
        if diagnostics.extraction_status == "extraction_failed":
            diagnostics.extraction_failure_reason = "No readable text was extracted from this file."
        return ExtractionResult(text=text, diagnostics=diagnostics)
    except Exception as exc:
        logger.exception("Document extraction failed for %s.", filename)
        diagnostics.extraction_status = "extraction_failed"
        diagnostics.extraction_failure_reason = _friendly_failure_reason(exc)
        return ExtractionResult(text="", diagnostics=diagnostics)


def document_likely_requires_ocr(filename: str, file_type: str | None, data: bytes) -> bool:
    if not OCR_ENABLED or not _is_pdf(filename, file_type):
        return False

    try:
        from pypdf import PdfReader
    except ImportError:
        return False

    try:
        reader = PdfReader(io.BytesIO(data))
        page_texts: list[str] = []
        image_page_count = 0
        for page in reader.pages:
            try:
                page_texts.append(page.extract_text() or "")
            except Exception:
                pass
            if _page_has_image_xobject(page):
                image_page_count += 1
        extracted_text = normalize_extracted_text("\n\n".join(page_texts))
        return len(extracted_text) < MIN_EXTRACTED_CHARACTERS and image_page_count > 0
    except Exception:
        logger.debug("Unable to preflight OCR requirement for %s.", filename, exc_info=True)
        return False


def _extract_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires the python-docx package.") from exc

    doc = docx.Document(io.BytesIO(data))
    return "\n\n".join(p.text for p in doc.paragraphs if p.text)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _extract_pdf(filename: str, file_type: str | None, data: bytes) -> ExtractionResult:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise RuntimeError("PDF parsing requires the pypdf package.") from exc

    reader = PdfReader(io.BytesIO(data))
    page_texts: list[str] = []
    image_page_count = 0
    text_page_count = 0

    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:
            page_text = ""
        normalized_page = normalize_extracted_text(page_text)
        if normalized_page:
            page_texts.append(normalized_page)
            text_page_count += 1
        if _page_has_image_xobject(page):
            image_page_count += 1

    extracted_text = normalize_extracted_text("\n\n".join(page_texts))
    diagnostics = ExtractionDiagnostics(
        page_count=len(reader.pages),
        extractor_used="pypdf",
        extracted_character_count=len(extracted_text),
        image_page_count=image_page_count,
        text_page_count=text_page_count,
        extraction_status="processing",
    )

    if diagnostics.extracted_character_count >= MIN_EXTRACTED_CHARACTERS:
        diagnostics.extraction_status = "searchable"
        return ExtractionResult(text=extracted_text, diagnostics=diagnostics)

    if image_page_count > 0:
        diagnostics.extraction_failure_reason = (
            "This PDF contains no readable text layer. OCR processing has been started."
        )
        if not OCR_ENABLED:
            diagnostics.extraction_status = "ocr_required"
            diagnostics.extraction_failure_reason = "This PDF contains no readable text layer. OCR is required."
            return ExtractionResult(text="", diagnostics=diagnostics)

        ocr_text, ocr_error = _ocr_pdf(data, diagnostics.page_count or 0)
        diagnostics.ocr_used = ocr_text is not None
        diagnostics.ocr_character_count = len(normalize_extracted_text(ocr_text or ""))
        if ocr_text and diagnostics.ocr_character_count >= MIN_EXTRACTED_CHARACTERS:
            diagnostics.extraction_status = "searchable"
            diagnostics.extraction_failure_reason = None
            return ExtractionResult(text=normalize_extracted_text(ocr_text), diagnostics=diagnostics)

        if ocr_error:
            diagnostics.extraction_status = "extraction_failed"
            diagnostics.extraction_failure_reason = ocr_error
            return ExtractionResult(text="", diagnostics=diagnostics)

        diagnostics.extraction_status = "ocr_required"
        diagnostics.extraction_failure_reason = "This PDF contains no readable text layer. OCR is required."
        return ExtractionResult(text="", diagnostics=diagnostics)

    diagnostics.extraction_status = "extraction_failed"
    diagnostics.extraction_failure_reason = "No readable text was extracted from this PDF."
    return ExtractionResult(text="", diagnostics=diagnostics)


def _ocr_pdf(data: bytes, page_count: int) -> tuple[str | None, str | None]:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError as exc:
        return None, f"OCR dependencies are unavailable: {exc.name or str(exc)}."

    try:
        pdf = pdfium.PdfDocument(data)
        texts: list[str] = []
        page_limit = min(len(pdf), MAX_OCR_PAGES if MAX_OCR_PAGES > 0 else len(pdf), page_count or len(pdf))
        for page_index in range(page_limit):
            page = pdf[page_index]
            bitmap = page.render(scale=2).to_pil()
            page_text = pytesseract.image_to_string(bitmap) or ""
            normalized = normalize_extracted_text(page_text)
            if normalized:
                texts.append(normalized)
        return normalize_extracted_text("\n\n".join(texts)), None
    except Exception as exc:
        logger.exception("OCR failed.")
        return None, _friendly_failure_reason(exc)


def _page_has_image_xobject(page: Any) -> bool:
    try:
        resources = page.get("/Resources") or {}
        xobjects = resources.get("/XObject") or {}
        if hasattr(xobjects, "get_object"):
            xobjects = xobjects.get_object()
        for item in xobjects.values():
            obj = item.get_object() if hasattr(item, "get_object") else item
            if obj.get("/Subtype") == "/Image":
                return True
    except Exception:
        return False
    return False


def _is_pdf(filename: str, file_type: str | None) -> bool:
    return filename.lower().endswith(".pdf") or "pdf" in (file_type or "").lower()


def _friendly_failure_reason(exc: Exception) -> str:
    message = str(exc).strip()
    if not message:
        return "Text extraction failed for this document."
    return message[:500]


def _optional_int(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _optional_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None
